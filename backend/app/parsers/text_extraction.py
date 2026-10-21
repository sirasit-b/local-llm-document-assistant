"""Text extraction for every supported document type.

Each extractor returns a page list. Formats without real pages (CSV, TXT) yield
a single logical page so that citations stay uniform across types.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path

import chardet
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from app.core.errors import ValidationError


@dataclass(frozen=True)
class ParsedPage:
    page_number: int
    content: str


@dataclass(frozen=True)
class ParsedDocument:
    pages: list[ParsedPage]
    #: True when the file is an image, or a PDF whose pages carry no text layer.
    needs_ocr: bool

    @property
    def char_count(self) -> int:
        return sum(len(p.content) for p in self.pages)


_WHITESPACE = re.compile(r"[ \t ]+")
_BLANK_LINES = re.compile(r"\n{3,}")


def clean_text(text: str) -> str:
    """Normalise whitespace while preserving paragraph structure."""
    normalised = text.replace("\r\n", "\n").replace("\r", "\n")
    normalised = _WHITESPACE.sub(" ", normalised)
    normalised = "\n".join(line.strip() for line in normalised.split("\n"))
    return _BLANK_LINES.sub("\n\n", normalised).strip()


def extract(path: Path, kind: str) -> ParsedDocument:
    extractor = _EXTRACTORS.get(kind)
    if extractor is None:
        raise ValidationError(f"ไม่รองรับไฟล์ประเภท {kind}")
    return extractor(path)


_THAI = re.compile(r"[฀-๿]")
_THAI_LATIN_SEAM = re.compile(r"[฀-๿][A-Za-z0-9=#)(]|[A-Za-z0-9=#)(][฀-๿]")

#: Above this share of Thai characters sitting directly against ASCII, the text
#: layer is assumed to be mis-mapped rather than genuinely mixed-language.
_GARBLED_SEAM_RATIO = 0.06


def is_garbled_thai(text: str) -> bool:
    """Detect a PDF text layer whose Thai characters decode to the wrong glyphs.

    Thai PDFs built from subset CID fonts often extract as Thai interleaved with
    stray ASCII ("แฟ้Iมสั่นี้กวั1าง"). Such a page carries text but no usable
    content, so it should go through OCR like a scan would.
    """
    thai_count = len(_THAI.findall(text))
    if thai_count < 20:
        return False
    seams = len(_THAI_LATIN_SEAM.findall(text))
    return seams / thai_count > _GARBLED_SEAM_RATIO


def _extract_pdf(path: Path) -> ParsedDocument:
    reader = PdfReader(str(path))
    pages: list[ParsedPage] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            raw = page.extract_text() or ""
        except Exception:  # a damaged page must not abort the whole document
            raw = ""
        pages.append(ParsedPage(page_number=index, content=clean_text(raw)))

    has_text = any(len(p.content) >= 40 for p in pages)
    garbled = bool(pages) and all(is_garbled_thai(p.content) for p in pages if p.content)
    return ParsedDocument(pages=pages, needs_ocr=not has_text or garbled)


def _extract_docx(path: Path) -> ParsedDocument:
    document = DocxDocument(str(path))
    blocks = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return _single_page("\n".join(blocks))


def _extract_xlsx(path: Path) -> ParsedDocument:
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    pages: list[ParsedPage] = []
    try:
        for index, sheet in enumerate(workbook.worksheets, start=1):
            lines = [f"[Sheet] {sheet.title}"]
            for row in sheet.iter_rows(values_only=True):
                values = ["" if v is None else str(v).strip() for v in row]
                if any(values):
                    lines.append(" | ".join(values))
            pages.append(ParsedPage(page_number=index, content=clean_text("\n".join(lines))))
    finally:
        workbook.close()
    return ParsedDocument(pages=pages or [ParsedPage(1, "")], needs_ocr=False)


def _extract_pptx(path: Path) -> ParsedDocument:
    presentation = Presentation(str(path))
    pages: list[ParsedPage] = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.extend(p.text for p in shape.text_frame.paragraphs if p.text.strip())
        pages.append(ParsedPage(page_number=index, content=clean_text("\n".join(lines))))
    return ParsedDocument(pages=pages or [ParsedPage(1, "")], needs_ocr=False)


def _extract_txt(path: Path) -> ParsedDocument:
    raw = path.read_bytes()
    encoding = chardet.detect(raw).get("encoding") or "utf-8"
    return _single_page(raw.decode(encoding, errors="replace"))


def _extract_csv(path: Path) -> ParsedDocument:
    raw = path.read_bytes()
    encoding = chardet.detect(raw).get("encoding") or "utf-8"
    text = raw.decode(encoding, errors="replace")
    reader = csv.reader(io.StringIO(text))
    lines = [" | ".join(cell.strip() for cell in row) for row in reader if any(row)]
    return _single_page("\n".join(lines))


def _extract_image(path: Path) -> ParsedDocument:
    # An image carries no text layer; OCR is the only source of content.
    return ParsedDocument(pages=[], needs_ocr=True)


def _single_page(text: str) -> ParsedDocument:
    return ParsedDocument(pages=[ParsedPage(1, clean_text(text))], needs_ocr=False)


_EXTRACTORS = {
    "pdf": _extract_pdf,
    "docx": _extract_docx,
    "xlsx": _extract_xlsx,
    "pptx": _extract_pptx,
    "txt": _extract_txt,
    "csv": _extract_csv,
    "image": _extract_image,
}
